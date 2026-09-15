# -*- coding: utf-8 -*-
"""
Genera un informe Word y PDF por empresa con el formato de referencia
(banner, tarjetas KPI y dos secciones), usando las observaciones consolidadas
por reporte_observaciones_empresa.py (su misma fuente de datos).

La agrupación de las reglas (Revisión General / Revisión Específica) proviene
del catálogo canónico definido en reporte_observaciones_empresa.CATALOGO.
"""

import os
import re
import sys
import json
import warnings
import pandas as pd
import psycopg2
from datetime import datetime
from collections import defaultdict

warnings.filterwarnings(
    "ignore",
    message="pandas only supports SQLAlchemy connectable",
    category=UserWarning,
)

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)
sys.path.insert(0, os.path.join(BASE_DIR, "scripts"))

import config
import reporte_observaciones_empresa as roe

REPORTE_DIR = config.REPORTE_DIR
DOCS_DIR = config.DOCS_DIR


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)


def set_table_borders(table, color="E2E8F0", sz="4", val="single"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(f'''
            <w:tblBorders {nsdecls("w")}>
                <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
                <w:insideV w:val="none"/>
                <w:left w:val="none"/>
                <w:right w:val="none"/>
            </w:tblBorders>
        ''')
        tblPr[0].append(borders)


def style_header_cell(cell, text, width_in_inches):
    cell.width = Inches(width_in_inches)
    set_cell_background(cell, "1E293B")
    set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.size = Pt(8.5)
    run.font.name = 'Arial'


def style_data_cell(cell, text, width_in_inches, align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, bg_hex=None):
    cell.width = Inches(width_in_inches)
    if bg_hex:
        set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.font.bold = bold
    run.font.size = Pt(8.5)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(51, 65, 85)


def add_section_header(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor(15, 23, 42)


# --- LECTURA DE DATOS ---

def _leer_total_registros(ruta):
    vacio = pd.DataFrame(columns=["empresa", "total_registros"])
    if not os.path.exists(ruta):
        return vacio
    df = pd.read_excel(ruta, sheet_name="analisis_por_empresa")
    if "empresa_name" not in df.columns or "total_registros" not in df.columns:
        return vacio
    return df[["empresa_name", "total_registros"]].rename(columns={"empresa_name": "empresa"})


def _clean(v):
    if v is None:
        return ""
    try:
        if pd.isna(v):
            return ""
    except (TypeError, ValueError):
        pass
    return str(v)


def cargar_total_registros():
    """Cuenta los registros EN_OPERACION por empresa en todas las tablas que
    poseen las columnas propietario_id y status (excluye tabla_unificada_nodos)."""
    totales = defaultdict(int)
    conn = psycopg2.connect(**config.DB_CONN)
    try:
        cur = conn.cursor()
        cur.execute(
            "SELECT table_name FROM information_schema.tables t "
            "WHERE table_schema='public' AND table_type='BASE TABLE' "
            "AND table_name <> 'tabla_unificada_nodos' "
            "AND EXISTS (SELECT 1 FROM information_schema.columns c "
            "            WHERE c.table_schema='public' AND c.table_name=t.table_name "
            "              AND c.column_name='propietario_id') "
            "AND EXISTS (SELECT 1 FROM information_schema.columns c "
            "            WHERE c.table_schema='public' AND c.table_name=t.table_name "
            "              AND c.column_name='status') "
            "ORDER BY table_name"
        )
        tablas = [r[0] for r in cur.fetchall()]
        cur.execute("SELECT id, name FROM api_empresa")
        empresa_map = {_clean(r[0]): r[1] for r in cur.fetchall()}
        cur.close()

        for t in tablas:
            df = pd.read_sql(
                f'SELECT propietario_id, COUNT(*) AS n FROM "{t}" '
                f"WHERE status='EN_OPERACION' GROUP BY propietario_id",
                conn,
            )
            for pid, n in zip(df["propietario_id"], df["n"]):
                nombre = empresa_map.get(_clean(pid), "")
                totales[nombre if nombre else "Sin empresa"] += int(n)
    finally:
        conn.close()

    return totales


# --- UNIVERSOS (total evaluado) POR REGLA ---
# Denominador del % error de cada fila del informe. Cada grupo de reglas se
# divide entre los registros que efectivamente evalúa esa regla (por empresa).

TABLAS_CONTENEDORES = [
    "api_subestacion", "api_patiosubestacion", "api_pano",
    "api_casaserviciosgenerales", "api_armario",
    "api_linea", "api_circuito", "api_tramo",
]

TABLAS_OOCC_ESTSE_SERV = ["api_oocc", "api_estructurasubestacion", "api_servidumbre"]

TABLAS_SUFIJO = ["api_circuito", "api_tramo"]

TABLAS_CONECTIVIDAD = ["api_circuito"]

TABLAS_HERENCIA = [
    "api_patiosubestacion", "api_casaserviciosgenerales", "api_elementocomunssee",
    "api_pano", "api_barra", "api_elementocomunpatiossee",
    "api_aislador", "api_bancocondensador", "api_compensadoractivo",
    "api_compensadorestaticoreactivo", "api_condensadoracoplamiento",
    "api_condensadorserie", "api_condensadorsincrono", "api_conexiontierra",
    "api_desconectador", "api_dispositivoreconexion", "api_elementocomunssggcaseta",
    "api_interruptor", "api_medidor", "api_mufa", "api_pararrayo", "api_reactor",
    "api_releproteccion", "api_tablero", "api_trampaonda",
    "api_transformadorcorriente", "api_transformadorpotencial",
    "api_transformadorssaa", "api_transformadorzigzag",
    "api_circuito",
]

TABLA_INSTALACION = {
    "Vano": "api_vano",
    "Torre": "api_torre",
    "Marcolinea": "api_marcolinea",
    "Tramo": "api_tramo",
    "Patio": "api_patiosubestacion",
}

UNIVERSOS_ESPECIFICOS = {
    "VANO_SIN_ACCESORIOS": ["api_vano"],
    "VANO_MULTIPLES_ACCESORIOS": ["api_vano"],
    "VANO_SIN_NODO1": ["api_vano"],
    "VANO_SIN_NODO2": ["api_vano"],
    "TORRE_SIN_OOCC": ["api_torre"],
    "TORRE_SIN_ACCESORIOS": ["api_torre"],
    "MARCO_SIN_OOCC": ["api_marcolinea"],
    "MARCO_SIN_ACCESORIOS": ["api_marcolinea"],
    "TRAMO_NODOS_IGUALES": ["api_tramo"],
    "TRAMO_NOMBRE_NODOS_NO_COINCIDEN": ["api_tramo"],
    "BARRA_TENSION_NO_COINCIDE": ["api_barra"],
    "PANO_TENSION_NO_COINCIDE": ["api_pano"],
    "PATIO_TENSION_NO_IDENTIFICADA": ["api_patiosubestacion"],
}


def _totales_excel(archivo):
    """Total evaluado por empresa desde la hoja 'analisis_por_empresa'."""
    df = _leer_total_registros(os.path.join(REPORTE_DIR, archivo))
    return dict(zip(df["empresa"], df["total_registros"])) if not df.empty else {}


def _contar_tablas_por_empresa(conn, tablas):
    """Suma los registros EN_OPERACION por empresa en las tablas indicadas."""
    cur = conn.cursor()
    cur.execute("SELECT id, name FROM api_empresa")
    empresa_map = {_clean(r[0]): r[1] for r in cur.fetchall()}
    cur.close()

    totales = defaultdict(int)
    for t in tablas:
        try:
            df = pd.read_sql(
                f'SELECT propietario_id, COUNT(*) AS n FROM "{t}" '
                f"WHERE status='EN_OPERACION' GROUP BY propietario_id",
                conn,
            )
        except Exception:
            continue
        for pid, n in zip(df["propietario_id"], df["n"]):
            nombre = empresa_map.get(_clean(pid), "")
            totales[nombre if nombre else "Sin empresa"] += int(n)
    return dict(totales)


def _tablas_prefijos():
    ruta = os.path.join(BASE_DIR, "utils", "prefijos.json")
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8") as f:
        data = json.load(f)
    return [e.get("tabla_psql", "") for e in data if e.get("tabla_psql")]


def _tablas_escritura(conn):
    cur = conn.cursor()
    cur.execute(
        "SELECT table_name FROM information_schema.tables t "
        "WHERE table_schema='public' AND table_type='BASE TABLE' "
        "AND table_name <> 'tabla_unificada_nodos' "
        "AND EXISTS (SELECT 1 FROM information_schema.columns c "
        "            WHERE c.table_name=t.table_name AND c.column_name='name') "
        "AND EXISTS (SELECT 1 FROM information_schema.columns c "
        "            WHERE c.table_name=t.table_name AND c.column_name='status') "
        "AND EXISTS (SELECT 1 FROM information_schema.columns c "
        "            WHERE c.table_name=t.table_name AND c.column_name='propietario_id') "
        "ORDER BY table_name"
    )
    tablas = [r[0] for r in cur.fetchall()]
    cur.close()
    return tablas


def cargar_universos():
    """Devuelve {clave: {empresa: total}} para cada fila del informe.

    Las claves son: subcategorías de Revisión General (nombre), instalaciones de
    Revisión Específica (nombre) y códigos de las reglas específicas.
    """
    conn = psycopg2.connect(**config.DB_CONN)
    try:
        u = {
            "Duplicidad por nombre": _totales_excel("analisis_duplicados.xlsx"),
            "Contenedores sin relacionamiento descendente": _contar_tablas_por_empresa(conn, TABLAS_CONTENEDORES),
            "Registros desconectados": _totales_excel("analisis_desconectados.xlsx"),
            "OOCC, Estructuras SE y Servidumbres sin conexión": _contar_tablas_por_empresa(conn, TABLAS_OOCC_ESTSE_SERV),
            "Revisión de prefijos": _contar_tablas_por_empresa(conn, _tablas_prefijos()),
            "Revisión de abreviatura de línea en torres": _totales_excel("revision_abreviatura_torre.xlsx"),
            "Revisión de herencia de nombre": _contar_tablas_por_empresa(conn, TABLAS_HERENCIA),
            "Revisión de sufijos": _contar_tablas_por_empresa(conn, TABLAS_SUFIJO),
            "Revisión de escritura general": _contar_tablas_por_empresa(conn, _tablas_escritura(conn)),
            "Conectividad de circuitos": _contar_tablas_por_empresa(conn, TABLAS_CONECTIVIDAD),
        }
        for instalacion, tabla in TABLA_INSTALACION.items():
            u[instalacion] = _contar_tablas_por_empresa(conn, [tabla])
        for codigo, tablas in UNIVERSOS_ESPECIFICOS.items():
            u[codigo] = _contar_tablas_por_empresa(conn, tablas)
        u["ABREVIATURA_LINEA_VANO"] = _totales_excel("revision_abreviatura_vano.xlsx")
        return u
    finally:
        conn.close()


def cargar_datos():
    """Consolida las observaciones (fuente única) y calcula los conteos por
    empresa y código, junto con los registros totales (base EN_OPERACION)."""
    df = roe.consolidar_observaciones()

    if not df.empty:
        conteo = df.groupby(["propietario_name", "observacion"]).size().reset_index(name="n")
    else:
        conteo = pd.DataFrame(columns=["propietario_name", "observacion", "n"])

    registros_totales = {k: int(v) for k, v in cargar_total_registros().items()}

    empresas = sorted(e for e in set(conteo["propietario_name"]) if e)

    return conteo, registros_totales, empresas


def _filas_empresa(conteo, empresa):
    if conteo.empty:
        return {}
    sub = conteo[conteo["propietario_name"] == empresa]
    return dict(zip(sub["observacion"], sub["n"]))


def _resumen_grupos(filas):
    """Devuelve (general, especifica) como listas de (subcategoria, cantidad)."""
    general = []
    especifica = []
    for entry in roe.CATALOGO:
        n = sum(filas.get(c, 0) for c in entry["codigos"])
        if entry["grupo"] == "General":
            general.append((entry["subcategoria"], n))
        else:
            especifica.append((entry["subcategoria"], n))
    return general, especifica


def _resumen_especifico(filas):
    """Agrupa la Revisión Específica por instalación (Patio/Vano/Torre/Marcolinea/Tramo).

    Devuelve una lista de diccionarios {instalacion, reglas, total}, incluyendo
    solo instalaciones con total > 0 y, dentro, solo reglas con afectados > 0.
    Cada regla es (codigo, etiqueta, cantidad).
    """
    resultado = []
    for grupo in roe.ESPECIFICAS_POR_INSTALACION:
        reglas = []
        total = 0
        for codigo, etiqueta in grupo["reglas"]:
            n = filas.get(codigo, 0)
            if n:
                reglas.append((codigo, etiqueta, n))
                total += n
        if total:
            resultado.append({
                "instalacion": grupo["instalacion"],
                "reglas": reglas,
                "total": total,
            })
    return resultado


def sanitizar_nombre(nombre):
    nombre = re.sub(r'[\\/:*?"<>|]', "_", str(nombre))
    nombre = nombre.strip().rstrip(".")
    return nombre[:120] or "Sin_nombre"


def _fmt(n):
    return f"{int(n):,}"


def _pct(e, t):
    return f"{e / t * 100:.2f}%" if t else "0.00%"


# --- GENERACIÓN DE WORD ---

def generar_documento(empresa, filas, registros_totales, universos, carpeta_word):
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)

    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(9.5)
    style.font.color.rgb = RGBColor(30, 41, 59)

    fecha = config.fecha_bd_str("%d/%m/%Y")

    total = registros_totales.get(empresa, 0)
    inconsistencias = sum(filas.values())

    general, _ = _resumen_grupos(filas)
    especifico = _resumen_especifico(filas)

    salud = round(max(0.0, 1 - inconsistencias / total) * 100, 2) if total else 100.0

    # 1. Banner
    header_table = doc.add_table(rows=1, cols=2)
    header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_cell_left = header_table.cell(0, 0)
    header_cell_right = header_table.cell(0, 1)
    set_cell_background(header_cell_left, "0F172A")
    set_cell_background(header_cell_right, "0F172A")
    set_cell_margins(header_cell_left, top=180, bottom=180, left=180, right=180)
    set_cell_margins(header_cell_right, top=180, bottom=180, left=180, right=180)
    header_cell_left.width = Inches(4.5)
    header_cell_right.width = Inches(2.5)

    p_title = header_cell_left.paragraphs[0]
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("Evaluación de Calidad de Datos")
    run_title.font.bold = True
    run_title.font.size = Pt(16)
    run_title.font.color.rgb = RGBColor(255, 255, 255)

    p_sub = header_cell_left.add_paragraph()
    p_sub.paragraph_format.space_after = Pt(0)
    run_sub = p_sub.add_run("BASE DE DATOS AUDITADA: BDATx / BDIT")
    run_sub.font.bold = True
    run_sub.font.size = Pt(9)
    run_sub.font.color.rgb = RGBColor(56, 189, 248)

    p_meta = header_cell_right.paragraphs[0]
    p_meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_meta.paragraph_format.space_after = Pt(0)
    p_meta.paragraph_format.line_spacing = 1.2
    run_meta = p_meta.add_run(
        f"Empresa: {empresa}\n"
        f"Fecha de Corte: {fecha}\n"
        f"Estado: Operativo e IRATx\n"
        f"Registros Totales: {_fmt(total)}"
    )
    run_meta.font.size = Pt(8)
    run_meta.font.color.rgb = RGBColor(148, 163, 184)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 2. Tarjetas KPI
    kpi_table = doc.add_table(rows=1, cols=3)
    kpi_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    kpi_data = [
        (f"{salud:.2f}%", "Salud Global", RGBColor(22, 163, 74)),
        (_fmt(total), "Registros Evaluados", RGBColor(15, 23, 42)),
        (_fmt(inconsistencias), "Inconsistencias", RGBColor(220, 38, 38)),
    ]

    for i, (val, lbl, color) in enumerate(kpi_data):
        cell = kpi_table.cell(0, i)
        cell.width = Inches(2.33)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=100, bottom=100, left=80, right=80)
        p_val = cell.paragraphs[0]
        p_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_val.paragraph_format.space_after = Pt(2)
        run_v = p_val.add_run(val)
        run_v.font.bold = True
        run_v.font.size = Pt(14)
        run_v.font.color.rgb = color
        p_lbl = cell.add_paragraph()
        p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_lbl.paragraph_format.space_after = Pt(0)
        run_l = p_lbl.add_run(lbl.upper())
        run_l.font.bold = True
        run_l.font.size = Pt(7)
        run_l.font.color.rgb = RGBColor(100, 116, 139)

    set_table_borders(kpi_table, color="E2E8F0")
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 3. Sección 1: Revisión General
    add_section_header(doc, "1. Revisión General")
    t1 = doc.add_table(rows=1 + len(general), cols=3)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t1)
    for col_idx, (h_text, w) in enumerate([("Criterio / Regla de Calidad", 4.7), ("Errores", 0.9), ("% Error", 1.0)]):
        style_header_cell(t1.cell(0, col_idx), h_text, w)

    for row_idx, (subcategoria, n) in enumerate(general, start=1):
        bg = "F8FAFC" if row_idx % 2 == 0 else "FFFFFF"
        denom = universos.get(subcategoria, {}).get(empresa, 0)
        style_data_cell(t1.cell(row_idx, 0), subcategoria, 4.7, bg_hex=bg)
        style_data_cell(t1.cell(row_idx, 1), _fmt(n), 0.9, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=bg)
        style_data_cell(t1.cell(row_idx, 2), _pct(n, denom), 1.0, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=bg)

    # 4. Sección 2: Revisión Específica (agrupada por instalación)
    add_section_header(doc, "2. Revisión Específica")
    n_filas_esp = sum(1 + len(g["reglas"]) for g in especifico)
    t3 = doc.add_table(rows=1 + n_filas_esp, cols=3)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t3)
    for col_idx, (h_text, w) in enumerate([("Regla Específica", 4.7), ("Afectados", 0.9), ("% Error", 1.0)]):
        style_header_cell(t3.cell(0, col_idx), h_text, w)

    row_idx = 1
    for g in especifico:
        bg_g = "E2E8F0"
        denom_grupo = universos.get(g["instalacion"], {}).get(empresa, 0)
        style_data_cell(t3.cell(row_idx, 0), g["instalacion"], 4.7, bold=True, bg_hex=bg_g)
        style_data_cell(t3.cell(row_idx, 1), _fmt(g["total"]), 0.9, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, bg_hex=bg_g)
        style_data_cell(t3.cell(row_idx, 2), _pct(g["total"], denom_grupo), 1.0, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, bg_hex=bg_g)
        row_idx += 1
        for j, (codigo, etiqueta, n) in enumerate(g["reglas"]):
            bg = "F8FAFC" if j % 2 == 0 else "FFFFFF"
            denom = universos.get(codigo, universos.get(g["instalacion"], {})).get(empresa, 0)
            style_data_cell(t3.cell(row_idx, 0), "    " + etiqueta, 4.7, bg_hex=bg)
            style_data_cell(t3.cell(row_idx, 1), _fmt(n), 0.9, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=bg)
            style_data_cell(t3.cell(row_idx, 2), _pct(n, denom), 1.0, align=WD_ALIGN_PARAGRAPH.CENTER, bg_hex=bg)
            row_idx += 1

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    ruta = os.path.join(carpeta_word, sanitizar_nombre(empresa) + ".docx")
    doc.save(ruta)


def convertir_a_pdf(carpeta_word, carpeta_pdf):
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        for nombre in os.listdir(carpeta_word):
            if not nombre.lower().endswith(".docx"):
                continue
            origen = os.path.join(carpeta_word, nombre)
            destino = os.path.join(carpeta_pdf, nombre[:-5] + ".pdf")
            documento = word.Documents.Open(origen)
            try:
                documento.SaveAs(destino, FileFormat=17)
            finally:
                documento.Close(False)
    finally:
        word.Quit()


def main():
    conteo, registros_totales, empresas = cargar_datos()
    universos = cargar_universos()

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    carpeta = os.path.join(DOCS_DIR, "ArchivoReporteEmpresa", ts + config.sufijo_carpeta_bd())
    carpeta_word = os.path.join(carpeta, "word")
    carpeta_pdf = os.path.join(carpeta, "pdf")
    os.makedirs(carpeta_word, exist_ok=True)
    os.makedirs(carpeta_pdf, exist_ok=True)

    for empresa in empresas:
        filas = _filas_empresa(conteo, empresa)
        generar_documento(empresa, filas, registros_totales, universos, carpeta_word)

    convertir_a_pdf(carpeta_word, carpeta_pdf)

    print(f"Se generaron {len(empresas)} documentos en {carpeta}")


if __name__ == "__main__":
    main()
